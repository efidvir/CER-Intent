import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Title
    title = lines[0].strip().replace('# ', '')
    h1 = doc.add_heading(title, 0)
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in lines[1:]:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('## '):
            doc.add_heading(line.replace('## ', ''), level=1)
        elif line.startswith('### '):
            doc.add_heading(line.replace('### ', ''), level=2)
        elif line.startswith('- '):
            p = doc.add_paragraph(line.replace('- ', ''), style='List Bullet')
        else:
            # Handle bold text in paragraphs
            p = doc.add_paragraph()
            parts = line.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 != 0:
                    run.bold = True

    doc.save(docx_path)
    print(f"Successfully created {docx_path}")

if __name__ == "__main__":
    md_file = "docs/intent_system_architecture.md"
    docx_file = "docs/intent_system_architecture.docx"
    convert_md_to_docx(md_file, docx_file)
