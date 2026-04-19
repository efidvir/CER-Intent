from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading('CER-Intent System Architecture', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('Modern 6G Transport Intent Management and Reconciliation Engine').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. Overview
    doc.add_heading('1. System Overview', level=1)
    doc.add_paragraph(
        'The CER-Intent platform is a Proof-of-Concept for Intent-Based Networking (IBN) '
        'tailored for Ceragon transport devices. It enables operators to define high-level '
        'service requirements (e.g., "Ensure 10 Gbps for URLLC slice") without needing '
        'to understand low-level CLI or MRMC script numbers.'
    )
    
    # 2. Smart Intent Reconciler
    doc.add_heading('2. Smart Intent Reconciliation Engine', level=1)
    doc.add_paragraph(
        'The core of the system is the Intent Architect, which has been upgraded to handle '
        'complex hardware constraints and state verification.'
    )
    
    doc.add_heading('2.1 Hardware-Aware Fallbacks', level=2)
    doc.add_paragraph(
        'When an intent configuration exceeds the physical capabilities of a target device, '
        'the system executes a strategy fallback. For example, if Space Diversity is requested '
        'on an IP-20C node (which lacks SD support), the Architect automatically reconciles '
        'this to a 1+1 HSB configuration. This ensures that the network is always in its '
        'best possible state rather than rejecting technical requests outright.'
    )
    
    doc.add_heading('2.2 State-Aware Optimization (No-Op)', level=2)
    doc.add_paragraph(
        'To reduce network control-plane load, the system performs a continuous state comparison. '
        'If a submitted intent corresponds to a state that is already active on the device, '
        'the system logs a "No-Action" event and skips the configuration cycle.'
    )
    
    # 3. Offline Resilience & Visualization
    doc.add_heading('3. Offline Resilience & Topology', level=1)
    doc.add_paragraph(
        'Designed for field deployment in restricted network environments, the dashboard '
        'operates with zero external CDN dependencies.'
    )
    
    doc.add_heading('3.1 Native SVG Visualizer', level=2)
    doc.add_paragraph(
        'The topology visualizer has been migrated from D3.js to a native SVG rendering engine. '
        'This allows for high-performance, interactive network plots (9 nodes, 10 links) '
        'to render perfectly even when internet access is blocked.'
    )
    
    doc.add_heading('3.2 Local Asset Management', level=2)
    doc.add_paragraph(
        'The Flask backend hosts all required assets locally, including Socket.IO libraries, '
        'fonts, and telemetry symbols, ensuring a premium user experience in air-gapped environments.'
    )

    # 4. Data Models
    doc.add_heading('4. Technical Data Models', level=1)
    doc.add_paragraph('The system utilizes a Python-based Pydantic schema for intent modeling and RFC-compliant YANG XML for device configuration.')
    
    doc.save('docs/intent_system_architecture_v2.docx')
    print('Document docs/intent_system_architecture_v2.docx created successfully.')

if __name__ == '__main__':
    create_doc()
